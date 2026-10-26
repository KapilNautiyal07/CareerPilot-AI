"""Text extraction from uploaded resume files (PDF, DOCX) with strict validation."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.errors import AppError
from app.services.untrusted import clean_text

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ALLOWED = {".pdf": PDF_MIME, ".docx": DOCX_MIME}
MAX_PDF_PAGES = 15
MAX_DOCX_UNCOMPRESSED = 30 * 1024 * 1024  # zip-bomb guard


class UnsupportedDocument(AppError):
    status_code = 415
    code = "unsupported_document"


def sniff_type(filename: str, data: bytes) -> str:
    """Determine the document type from magic bytes; the extension must agree."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise UnsupportedDocument("Only PDF and DOCX resumes are supported")
    if data.startswith(b"%PDF-"):
        actual = ".pdf"
    elif data.startswith(b"PK\x03\x04"):
        actual = ".docx"
    else:
        raise UnsupportedDocument("File content is not a valid PDF or DOCX document")
    if actual != ext:
        raise UnsupportedDocument("File extension does not match its content")
    return ALLOWED[actual]


def extract_text(data: bytes, mime: str) -> str:
    try:
        if mime == PDF_MIME:
            return _pdf_text(data)
        if mime == DOCX_MIME:
            return _docx_text(data)
    except UnsupportedDocument:
        raise
    except (PdfReadError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise UnsupportedDocument(f"Could not read document: {exc.__class__.__name__}") from exc
    raise UnsupportedDocument("Unsupported document type")


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise UnsupportedDocument("Encrypted PDFs are not supported")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise UnsupportedDocument(f"Resume has more than {MAX_PDF_PAGES} pages")
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return _finish(text)


def _docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if sum(i.file_size for i in zf.infolist()) > MAX_DOCX_UNCOMPRESSED:
            raise UnsupportedDocument("Document is too large when decompressed")
        if "word/document.xml" not in zf.namelist():
            raise UnsupportedDocument("Not a Word document")
    doc = Document(io.BytesIO(data))
    lines: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return _finish("\n".join(lines))


def _finish(text: str) -> str:
    text = clean_text(text)
    if len(text) < 50:
        raise UnsupportedDocument(
            "Very little text could be extracted. Scanned/image-only resumes are not supported; "
            "please upload a text-based file."
        )
    return text[:60_000]
